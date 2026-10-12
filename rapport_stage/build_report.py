"""Génère le rapport de stage technicien (DOCX) de Taher Yahia.

Usage :
    python build_figures.py   # génère les figures
    python build_report.py    # génère Rapport_Stage_Taher_Yahia.docx

Le document utilise des champs Word (SEQ pour les légendes, REF pour les renvois,
TOC pour le sommaire et les listes). Word demande de mettre à jour les champs à
l'ouverture (paramètre updateFields) ; sinon : Ctrl+A puis F9.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
FIG = HERE / "figures"
OUTPUT = HERE / "Rapport_Stage_Taher_Yahia.docx"
FONT = "Times New Roman"

# Numérotation attendue (vérifiée à la génération par assert)
FIG_ORDER = ["fig_arch", "fig_model"]
TAB_ORDER = ["tab_bf", "tab_stack", "tab_chrono", "tab_regles",
             "tab_tests", "tab_limites", "tab_comp", "tab_end"]
LABELS = {}
for i, k in enumerate(FIG_ORDER, 1):
    LABELS[k] = ("Figure", i)
for i, k in enumerate(TAB_ORDER, 1):
    LABELS[k] = ("Tableau", i)

doc = Document()
_counters = {"Figure": 0, "Tableau": 0}
_bm_id = [100]


# ----------------------------------------------------------------- utilitaires
def fr(text: str) -> str:
    """Espaces insécables avant ; : ? ! % (typographie française)."""
    return re.sub(r" (?=[;:?!%])", "\u00a0", text)


def _rfonts(rpr, font=FONT):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(att)) is not None:
            del rf.attrib[qn(att)]
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(att), font)


def style_font(style, size, bold=False, italic=False, color_black=True):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if color_black:
        style.font.color.rgb = RGBColor(0, 0, 0)
    _rfonts(style.element.get_or_add_rPr())


def add_run(par, text, bold=False, italic=False, size=None, font=None):
    r = par.add_run(fr(text))
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    _rfonts(r._r.get_or_add_rPr(), font or FONT)
    return r


def add_field(par, instr, result="", size=None, bold=False):
    def fld(kind):
        r = par.add_run()
        if size:
            r.font.size = Pt(size)
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
        return r
    fld("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fld("separate")
    add_run(par, result, bold=bold, size=size)
    fld("end")


def bookmark(par, name, start=True):
    el = OxmlElement("w:bookmarkStart" if start else "w:bookmarkEnd")
    if start:
        _bm_id[0] += 1
        el.set(qn("w:id"), str(_bm_id[0]))
        el.set(qn("w:name"), name)
    else:
        el.set(qn("w:id"), str(_bm_id[0]))
    par._p.append(el)


def para(text, style=None, align=None, space_after=None, keep_next=False):
    """Paragraphe avec mini-balisage : **gras** et {ref:cle} (renvoi REF)."""
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if keep_next:
        p.paragraph_format.keep_with_next = True
    for part in re.split(r"((?<![\w/*])\*\*(?=\S).+?(?<=\S)\*\*(?![\w*])|\{ref:\w+\})", text):
        if not part:
            continue
        if part.startswith("**"):
            add_run(p, part[2:-2], bold=True)
        elif part.startswith("{ref:"):
            key = part[5:-1]
            label, num = LABELS[key]
            add_field(p, f"REF ref_{key} \\h", f"{label} {num}")
        else:
            add_run(p, part)
    return p


def bullets(items):
    for it in items:
        para(it, style="List Bullet", space_after=2)


def h1(text):
    doc.add_paragraph(text, style="Heading 1")


def h2(text):
    doc.add_paragraph(text, style="Heading 2")


def caption(kind, key, text, above=False):
    label, num = LABELS[key]
    assert label == kind and _counters.get(kind, 0) + 1 == num, (key, kind, num)
    _counters[kind] = num
    p = doc.add_paragraph(style="Caption")
    if above:
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bookmark(p, f"ref_{key}", start=True)
    add_run(p, f"{kind} ", bold=True, size=10)
    add_field(p, f"SEQ {kind} \\* ARABIC", str(num), size=10, bold=True)
    bookmark(p, f"ref_{key}", start=False)
    add_run(p, f" — {text}", size=10)
    return p


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table(headers, rows, widths, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.width = Cm(widths[i])
        shade(c, "D9D9D9")
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, h, bold=True, size=size)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Cm(widths[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_run(p, val, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, width_cm=15.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def code(text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, line, size=9.5, font="Courier New")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def page_break():
    doc.add_paragraph().add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE)


# ----------------------------------------------------------------- mise en page
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.5)
sec.right_margin = Cm(2.5)
sec.left_margin = Cm(3.0)
sec.header_distance = Cm(1.0)
sec.footer_distance = Cm(1.0)
sec.different_first_page_header_footer = True
TEXT_WIDTH_CM = 21.0 - 3.0 - 2.5

normal = doc.styles["Normal"]
style_font(normal, 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

style_font(doc.styles["Heading 1"], 14, bold=True)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(9)
doc.styles["Heading 1"].paragraph_format.keep_with_next = True
doc.styles["Heading 1"].paragraph_format.line_spacing = 1.15
style_font(doc.styles["Heading 2"], 12, bold=True)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)
doc.styles["Heading 2"].paragraph_format.keep_with_next = True
doc.styles["Heading 2"].paragraph_format.line_spacing = 1.15
style_font(doc.styles["Caption"], 10, bold=False, italic=False)
doc.styles["Caption"].paragraph_format.line_spacing = 1.0
doc.styles["Caption"].paragraph_format.space_before = Pt(4)
doc.styles["Caption"].paragraph_format.space_after = Pt(8)
style_font(doc.styles["List Bullet"], 12)

# En-tête et pied de page (hors page de garde)
hp = sec.header.paragraphs[0]
hp.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
add_run(hp, "Mini-API Assurance — VERMEG\tTaher Yahia", size=9)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_field(fp, "PAGE", "1", size=10)

# Paramètre : Word met à jour les champs (sommaire, légendes, renvois) à l'ouverture
settings = doc.settings.element
upd = OxmlElement("w:updateFields")
upd.set(qn("w:val"), "true")
settings.append(upd)

# ----------------------------------------------------------------- page de garde
def centered(text, size=12, bold=False, space_after=4, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    add_run(p, text, bold=bold, size=size, italic=italic)
    return p


centered("Institut Supérieur d'Informatique et de Multimédia de Gabès (ISIMG)", 12, True, 2)
centered("Université de Gabès", 12, False, 18)

lt = doc.add_table(rows=1, cols=2)
lt.alignment = WD_TABLE_ALIGNMENT.CENTER
lt.autofit = False
for i, label in enumerate(["Emplacement du logo ISIMG\n(à insérer)", "Emplacement du logo VERMEG\n(à insérer)"]):
    c = lt.rows[0].cells[i]
    c.width = Cm(7.5)
    shade(c, "F2F2F2")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(40)
    add_run(p, label, size=9, italic=True)
doc.add_paragraph().paragraph_format.space_after = Pt(14)

centered("RAPPORT DE STAGE TECHNICIEN", 16, True, 10)
centered("Conception et réalisation d'une mini-API REST de gestion d'assurance", 14, True, 28)

for line in [
    "**Réalisé par :** Taher Yahia",
    "**Formation :** génie logiciel, FGL1 au moment du stage",
    "**Organisme d'accueil :** VERMEG",
    "**Encadrante en entreprise :** Mme Faten Kardous",
    "**Période de stage :** du 1er au 31 juillet 2026",
    "**Année universitaire :** 2025/2026",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    for part in re.split(r"(\*\*.+?\*\*)", line):
        if part.startswith("**"):
            add_run(p, part[2:-2], bold=True)
        elif part:
            add_run(p, part)

centered("Version 1 — projet de rapport (octobre 2026)", 10, False, 0, italic=True)
page_break()

# ----------------------------------------------------------------- sommaire et listes
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
add_run(p, "Sommaire", bold=True, size=14)
p = doc.add_paragraph()
add_field(p, 'TOC \\o "1-2" \\h \\z \\u',
          "Sommaire à actualiser : clic droit sur la table, puis Mettre à jour les champs.")
p.paragraph_format.line_spacing = 1.15
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
add_run(p, "Liste des figures", bold=True, size=14)
p = doc.add_paragraph()
add_field(p, 'TOC \\h \\z \\c "Figure"', "Liste à actualiser (Mettre à jour les champs).")
p.paragraph_format.line_spacing = 1.15
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
add_run(p, "Liste des tableaux", bold=True, size=14)
p = doc.add_paragraph()
add_field(p, 'TOC \\h \\z \\c "Tableau"', "Liste à actualiser (Mettre à jour les champs).")
p.paragraph_format.line_spacing = 1.15
page_break()

# ----------------------------------------------------------------- 1. introduction
h1("1. Introduction générale")
para("Ce rapport présente le stage technicien que j'ai effectué chez VERMEG du 1er au 31 juillet 2026. "
     "Il s'inscrit dans ma formation en génie logiciel à l'Institut Supérieur d'Informatique et de Multimédia "
     "de Gabès (ISIMG) : j'étais alors inscrit en FGL1 (désignation à confirmer selon le guide). Le stage a été encadré "
     "en entreprise par Mme Faten Kardous.")
para("Le sujet consistait à concevoir et réaliser une mini-API REST de gestion d'assurance. L'application "
     "gère trois types d'objets : les clients, les contrats d'assurance et les sinistres déclarés sur un contrat. "
     "Le travail a porté sur le service serveur : modèle de données, interface REST, authentification par jetons, "
     "règles métier, documentation et conteneurisation. Une interface web Angular a également été développée "
     "au cours du stage pour consommer cette API.")
para("La question technique qui guide le rapport est la suivante : comment exposer, de façon sécurisée et "
     "cohérente, les opérations de gestion d'un portefeuille d'assurance, en faisant respecter les règles métier "
     "par le serveur plutôt que par le client ?")
para("Le chapitre 2 présente VERMEG et le cadre du stage, le chapitre 3 la conception, le chapitre 4 la réalisation et les vérifications, le chapitre 5 les limites et le chapitre 6 le bilan. Le rapport distingue ce qui est constaté dans le dépôt de ce qui reste à confirmer.")

# ----------------------------------------------------------------- 2. VERMEG et cadre
h1("2. Présentation de VERMEG et contexte du stage")
h2("2.1 VERMEG")
para("VERMEG est un éditeur de solutions logicielles pour le secteur financier. Un communiqué de Nordic Capital publié en novembre 2024 indique que l'entreprise a été fondée en 1993, qu'elle a son siège à Amsterdam et qu'elle fournit des solutions aux secteurs de la banque, de l'assurance et de la gestion de patrimoine [1]. Deux profils d'entreprise publics indiquent la même date de fondation et le même siège [2], [3]. Ces sources ne sont pas des documents officiels de VERMEG : l'organisation interne et les produits de l'entreprise ne sont pas décrits ici, faute de source officielle consultée.")
h2("2.2 Cadre du stage")
para("Le stage s'est déroulé du 1er au 31 juillet 2026 inclus, sous la forme d'un stage technicien [à confirmer selon le guide]. Pendant la "
     "première semaine, j'ai installé et configuré l'environnement de développement, étudié le contexte du projet "
     "et la logique métier de l'assurance, puis réfléchi à la modélisation des données : entités, attributs et "
     "relations. Les propositions de modélisation ont été relues avec l'encadrante, puis un modèle a été retenu. "
     "La réalisation a suivi ce modèle. Le dépôt GitHub "
     "du projet a été créé pendant cette phase, le 6 juillet 2026 [4].")
h2("2.3 Objectifs")
para("L'objectif principal était de réaliser un service REST qui gère les clients, les contrats et les sinistres, "
     "et dont le serveur garantit la cohérence des données. Les objectifs secondaires étaient de sécuriser l'accès "
     "aux ressources par un jeton d'authentification, de documenter l'interface de façon interactive, de rendre "
     "l'environnement d'exécution reproductible et de proposer une interface web capable de consommer l'API.")
h2("2.4 Démarche de travail")
para("Le développement a été conduit par incréments, chacun enregistré dans un commit distinct. La branche principale comporte 36 commits, du 5 au 30 juillet 2026 ; cet historique est la principale trace de la chronologie du chapitre 4.")

# ----------------------------------------------------------------- 3. besoin et conception
h1("3. Étude du besoin et conception")
h2("3.1 Contexte métier")
para("Un système d'assurance manipule trois notions liées. Un **client** est une personne identifiée par son état "
     "civil et par un identifiant national, le CIN, que l'application valide sur huit chiffres. Un **contrat** relie "
     "un client à une couverture : il porte un numéro de police, un type (AUTO, HOME, HEALTH ou LIFE), une période "
     "de validité, un montant couvert, une prime et un statut (ACTIVE, EXPIRED ou TERMINATED). Un **sinistre** est un "
     "événement dommageable déclaré sur un contrat : il porte un numéro, une description, une date de survenance, une "
     "date de déclaration, un montant estimé, un remboursement et un statut.")
para("Trois invariants structurent le domaine : un sinistre appartient à un contrat, un contrat appartient à un "
     "client, et un sinistre ne peut être déclaré que sur un contrat actif, pour une date comprise dans la période "
     "de couverture. Le besoin central était que le serveur garantisse ces invariants, indépendamment du client qui l'appelle.")
h2("3.2 Besoins fonctionnels")
para("Le tableau {ref:tab_bf} recense les besoins fonctionnels et leur état, tel que constaté dans le code source.", keep_next=True)
caption("Tableau", "tab_bf", "Besoins fonctionnels et état de réalisation dans le dépôt", above=True)
table(["Réf.", "Besoin", "État constaté"], [
    ["BF1", "Créer un compte utilisateur avec un rôle", "Réalisé, par inscription publique (voir 4.3 et 5.2)"],
    ["BF2", "S'authentifier et obtenir un jeton d'accès", "Réalisé (jeton valable 24 heures)"],
    ["BF3", "Créer, consulter et modifier un client", "Réalisé"],
    ["BF4", "Archiver ou supprimer un client", "Partiel : archivage logique, sans contrôle des contrats liés"],
    ["BF5", "Créer un contrat pour un client existant", "Réalisé"],
    ["BF6", "Consulter les contrats (liste, détail, par client)", "Réalisé"],
    ["BF7", "Mettre à jour un contrat (montants, date de fin, statut)", "Réalisé, sans contrôle de cohérence des dates"],
    ["BF8", "Déclarer un sinistre sur un contrat actif", "Réalisé ; justificatif facultatif"],
    ["BF9", "Consulter les sinistres (par contrat, liste, détail)", "Réalisé"],
    ["BF10", "Instruire un sinistre (transitions, remboursement)", "Non réalisé : les statuts existent, aucune transition n'est codée"],
    ["BF11", "Gérer les comptes utilisateurs (ADMIN)", "Partiel : liste réservée à l'ADMIN ; création possible sans authentification"],
], [1.4, 6.6, 7.5])
h2("3.3 Besoins non fonctionnels")
para("Cinq exigences non fonctionnelles guident la conception : **sécurité** (aucune ressource métier sans jeton), **intégrité** (unicité et non-nullité déclarées dans le mapping JPA), **validation** des données d'entrée, **documentation** OpenAPI et **portabilité** (démarrage par Docker Compose). Le chapitre 4 indique dans quelle mesure elles sont tenues.")
h2("3.4 Architecture de la solution")
para("L'application suit une architecture en couches, représentée à la figure {ref:fig_arch}. Une requête HTTP "
     "traverse d'abord le filtre JwtAuthFilter, qui lit le jeton et renseigne le contexte de sécurité. Le contrôleur "
     "REST reçoit des objets de transfert (DTO) validés et délègue le traitement à un service. Le service applique les "
     "règles métier, convertit les entités en DTO de réponse et s'appuie sur un repository Spring Data JPA pour lire "
     "ou écrire dans PostgreSQL. Les exceptions métier sont converties en réponses JSON par un gestionnaire global. "
     "Le frontend Angular, lancé en développement sur le port 4200 (ng serve), appelle l'API sur le port 8080 : cette origine différente impose "
     "une configuration CORS.")
figure(FIG / "fig1_architecture.png", 14.5)
caption("Figure", "fig_arch", "Architecture en couches de la mini-API (d'après la structure du dépôt)")

h2("3.5 Modèle de données")
para("Le modèle comporte quatre entités persistantes, représentées à la figure {ref:fig_model}. Plusieurs choix de "
     "modélisation sont visibles dans le code :", keep_next=True)
bullets([
    "les montants sont typés BigDecimal et non double, afin d'éviter les erreurs d'arrondi sur des valeurs monétaires ;",
    "les énumérations sont enregistrées sous forme de texte (@Enumerated(STRING)) : le stockage ne dépend pas de la "
    "position des constantes dans l'énumération ;",
    "les associations ManyToOne vers Client et Contract sont chargées à la demande (LAZY), et l'option open-in-view "
    "est désactivée dans la configuration ;",
    "l'utilisateur est mappé sur la table app_user, car le commentaire de la classe signale que « user » peut entrer "
    "en conflit avec un mot réservé ;",
    "la suppression d'un client est logique : le champ active passe à false et le client disparaît des listes.",
])
figure(FIG / "fig2_modele_donnees.png", 14.5)
caption("Figure", "fig_model", "Modèle de données : entités principales et relations")
h2("3.6 Choix techniques")
para("Le tableau {ref:tab_stack} présente les composants retenus, leur version telle que déclarée dans le dépôt et "
     "leur rôle dans la solution.", keep_next=True)
caption("Tableau", "tab_stack", "Composants techniques, versions déclarées et rôle", above=True)
table(["Composant", "Version déclarée", "Rôle dans la solution"], [
    ["Java", "17 (pom.xml)", "Langage du backend ; records pour les DTO"],
    ["Spring Boot", "3.5.16 (pom.xml)", "Démarrage, configuration, serveur HTTP embarqué"],
    ["Spring Data JPA / Hibernate", "Gérés par Spring Boot", "Persistance objet-relationnelle, repositories"],
    ["PostgreSQL", "16 (docker-compose.yml)", "Base relationnelle, contraintes d'intégrité"],
    ["Spring Security et JJWT", "JJWT 0.12.6", "Authentification par jeton signé (HMAC)"],
    ["Bean Validation", "Géré par Spring Boot", "Validation déclarative des DTO"],
    ["springdoc-openapi", "2.7.0", "Spécification OpenAPI et interface Swagger UI"],
    ["Docker et Compose", "maven:3.9.9-eclipse-temurin-17, eclipse-temurin:17-jre, postgres:16", "Construction et exécution reproductibles"],
    ["Angular", "21.2 (package.json)", "Interface web consommant l'API"],
    ["Chart.js", "4.5.1 (package.json)", "Graphique du tableau de bord"],
], [4.2, 4.9, 6.4])


# ----------------------------------------------------------------- 4. réalisation
h1("4. Réalisation technique et validation")
h2("4.1 Organisation du code")
para("Le code backend est réparti en paquetages (domain, dto, mapper, repository, service, controller, security, config, exception). Les dépendances sont injectées par constructeur et les DTO sont des records Java immuables.")
para("Les contrôles d'une déclaration de sinistre illustrent l'emplacement des règles métier. Ils sont placés dans "
     "la couche service, avant toute écriture en base :", keep_next=True)
code("""
if (contract.getStatus() != ContractStatus.ACTIVE) {
    throw new BusinessRuleException("Cannot file a claim: Contract is not ACTIVE");
}
if (dto.claimDate().isBefore(contract.getStartDate())
        || dto.claimDate().isAfter(contract.getEndDate())) {
    throw new BusinessRuleException("Claim date must be within the contract coverage period");
}
""")
h2("4.2 Chronologie du développement")
para("Le tableau {ref:tab_chrono} regroupe les travaux par périodes, à partir des dates et des messages des commits "
     "de la branche principale.", keep_next=True)
caption("Tableau", "tab_chrono", "Chronologie des travaux d'après l'historique Git", above=True)
table(["Période", "Travaux observés dans l'historique", "Commits"], [
    ["5 au 8 juillet", "Initialisation du projet Spring Boot ; configuration de PostgreSQL ; entités Client, Contrat et Sinistre, CRUD des clients ; "
     " gestion globale des erreurs ; OpenAPI ; authentification JWT et Spring Security ; "
     "données de démonstration ; README ; passage du code en anglais", "14"],
    ["13 au 15 juillet", "Enrichissement du modèle (client, contrat, sinistre) ; règles métier sur les sinistres ; "
     "mise à jour des clients ; suppression de client avec contrôle des contrats ; mise à jour des contrats ; "
     "liste des sinistres par contrat ; premier test unitaire", "9"],
    ["21 au 22 juillet", "Interface Angular : service d'authentification, intercepteur JWT, garde de route, liste et "
     "création de clients", "6"],
    ["29 au 30 juillet", "Barre de navigation et tableau de bord ; règles métier et archivage logique des clients ; écrans "
     "contrats et sinistres ; formulaire de déclaration de sinistre ; contrôle d'accès par rôle côté backend ; "
     "gestion des utilisateurs", "7"],
], [3.0, 11.0, 2.0])
para("Aucun commit n'apparaît entre le 9 et le 12 juillet, ni entre le 16 et le 20, ni entre le 23 et le 28 juillet. Les dates marquent l'enregistrement des modifications, non le moment où chaque tâche a été accomplie.")
h2("4.3 Authentification et contrôle d'accès")
para("L'authentification repose sur des jetons JWT signés par une clé HMAC. Le point d'entrée POST /api/auth/login "
     "vérifie le couple identifiant et mot de passe, puis renvoie un jeton valable 24 heures. Les mots de passe sont "
     "hachés par BCrypt avant d'être enregistrés. Pour chaque requête suivante, le filtre JwtAuthFilter lit l'en-tête "
     "Authorization, valide le jeton et renseigne le contexte de sécurité. La session est STATELESS : le serveur ne "
     "conserve aucun état entre deux requêtes. La protection CSRF est désactivée, le jeton étant transmis dans un "
     "en-tête et non dans un cookie.")
para("La configuration SecurityConfig applique les règles suivantes, par chemin : les routes /api/auth/**, la "
     "documentation Swagger et OpenAPI sont publiques ; /api/users/** est réservé au rôle ADMIN ; /api/clients/**, "
     "/api/contracts/** et /api/claims/** sont accessibles aux rôles ADMIN et AGENT ; toute autre route exige une "
     "authentification.")
para("Des annotations @PreAuthorize figurent également sur plusieurs méthodes des contrôleurs. L'annotation "
     "@EnableMethodSecurity n'est toutefois déclarée nulle part : ces annotations ne sont donc pas évaluées. Un compte "
     "AGENT peut ainsi appeler DELETE /api/clients/{id} et PUT /api/contracts/{id}, alors que ces méthodes sont "
     "annotées pour le seul rôle ADMIN (section 5.2).")
para("Enfin, la route POST /api/auth/register est publique et accepte un champ role dans le corps de la requête. "
     "D'après la lecture du code, un appel sans authentification peut donc créer un compte ADMIN. Ce point n'a pas "
     "été testé à l'exécution, mais il constitue la limite la plus importante du projet (section 5.2).")
h2("4.4 Gestion des clients, des contrats et des sinistres")
para("L'API expose 17 points d'accès, détaillés en annexe A. Le serveur y fait respecter les règles du tableau "
     "{ref:tab_regles}. Ces contrôles sont exécutés dans la couche service, indépendamment du client qui appelle l'API.",
     keep_next=True)
caption("Tableau", "tab_regles", "Règles métier et emplacement dans le code", above=True)
table(["Réf.", "Règle", "Emplacement", "Comportement"], [
    ["RM1", "Date de fin de contrat postérieure à la date de début (à la création)", "ContractService", "Erreur 400"],
    ["RM2", "Contrat créé au statut ACTIVE, avec numéro de police généré par le serveur", "ContractService", "Valeurs imposées"],
    ["RM3", "Statut de contrat valide lors de la mise à jour", "ContractService", "Erreur 400"],
    ["RM4", "Sinistre refusé si le contrat n'est pas ACTIVE", "ClaimService", "Erreur 400"],
    ["RM5", "Date de survenance comprise dans la période de couverture", "ClaimService", "Erreur 400"],
    ["RM6", "Sinistre créé en SUBMITTED, remboursement à zéro, date de déclaration fixée par le serveur", "ClaimService", "Valeurs imposées"],
    ["RM7", "CIN unique de 8 chiffres et téléphone de 8 chiffres (à la création)", "ClientCreateDto, ClientService", "Erreur 400"],
    ["RM8", "Client âgé d'au moins 18 ans à la création", "ClientService", "Erreur 400"],
    ["RM9", "Client archivé non modifiable ; suppression logique", "ClientService", "Erreur 400 ; archivage"],
    ["RM10", "Nom d'utilisateur et courriel uniques", "UserService", "Erreur 400"],
    ["RM11", "Client lié à un contrat non supprimable", "Non appliquée (existsByClientId non appelée)", "Aucun"],
], [1.4, 7.4, 4.4, 2.8])
para("La règle RM11 figure dans le tableau pour être explicite : elle a été codée dans une version antérieure, "
     "puis n'a plus été appliquée (section 5.1).")
para("Les sinistres peuvent être accompagnés d'un justificatif envoyé en multipart/form-data. Le fichier est "
     "enregistré dans le dossier uploads/ sous un nom UUID, son extension d'origine est conservée, et son chemin est "
     "stocké dans l'entité Claim. Aucun contrôle de type ni de taille n'est effectué (section 5.2).")
h2("4.5 Documentation OpenAPI")
para("La bibliothèque springdoc-openapi génère la spécification OpenAPI à partir des contrôleurs et expose "
     "l'interface Swagger UI à l'adresse /swagger-ui/index.html. La classe OpenApiConfig déclare un schéma de "
     "sécurité bearerAuth : le bouton Authorize permet de saisir le jeton obtenu à la connexion, puis de tester les "
     "routes protégées depuis le navigateur. ")
h2("4.6 Interface web Angular")
para("L'interface compte 11 écrans : connexion, tableau de bord, liste des clients, détail d'un client, création d'un "
     "client, liste des contrats, détail d'un contrat, liste des sinistres, détail d'un sinistre, formulaire de "
     "déclaration de sinistre et gestion des employés. Le service d'authentification conserve le jeton dans le "
     "localStorage. Un intercepteur HTTP ajoute l'en-tête Authorization à chaque requête. Deux gardes de route "
     "complètent le dispositif : authGuard protège les pages, et roleGuard réserve la gestion des employés au rôle ADMIN.")
para("Cette interface présente des limites. L'URL de l'API est écrite en dur (http://localhost:8080) dans plusieurs "
     "services. L'intercepteur ne traite pas les réponses 401. Le jeton, stocké dans le localStorage, est accessible à "
     "tout script exécuté sur la page. Depuis l'interface, aucun client ni aucun contrat ne peut être modifié ni "
     "supprimé. La page de gestion des employés crée les comptes par la route publique d'inscription (section 5.2).")
h2("4.7 Conteneurisation")
para("Le Dockerfile construit l'application en deux étapes : une image Maven avec JDK 17 compile le projet et produit "
     "l'archive exécutable, puis une image JRE 17 n'embarque que cette archive. Le fichier docker-compose.yml décrit "
     "deux services : db (PostgreSQL 16, avec un contrôle de santé pg_isready et un volume nommé pgdata) et api, qui "
     "ne démarre qu'une fois la base déclarée saine. La configuration est injectée par variables d'environnement, "
     "avec des valeurs par défaut dans application.yaml.")
para("Cette configuration présente des limites. Le conteneur s'exécute en root. Les identifiants figurent en clair dans "
     "le fichier de composition. Les tests sont ignorés lors de la construction de l'image. Enfin, le dossier uploads/ "
     "n'est monté dans aucun volume : les justificatifs sont perdus si le conteneur est recréé.")
h2("4.8 Tests et vérifications")
para("Le tableau {ref:tab_tests} recense les tests présents dans le dépôt.", keep_next=True)
caption("Tableau", "tab_tests", "Tests présents dans le dépôt", above=True)
table(["Test", "Type", "Dépendances", "Ce qu'il vérifie"], [
    ["ClaimServiceTest : refus d'un sinistre sur un contrat TERMINATED", "Unitaire (JUnit 5, Mockito)", "Aucune base", "Règle RM4"],
    ["MiniApiAssuranceApplicationTests : chargement du contexte", "Démarrage Spring", "PostgreSQL accessible", "Configuration et démarrage"],
    ["Tests frontend (9 fichiers .spec.ts)", "Générés par Angular CLI", "Environnement de test Angular", "Instanciation des composants ; non exécutés ici"],
], [5.4, 3.6, 3.4, 3.6])
para("Le dépôt contient donc un test unitaire ciblant une règle métier parmi les dix appliquées. Il n'existe ni test "
     "de contrôleur, ni test de sécurité, ni test d'intégration sur une base, ni mesure de couverture. Les tests "
     "frontend ne vérifient aucun scénario métier.")
para("Les vérifications réalisées pour ce rapport sont une lecture du code backend (contrôleurs, services, DTO, entités, configuration de sécurité, gestion des erreurs, tests et fichiers de configuration), d'une partie du code Angular, et une analyse de l'historique Git via l'API GitHub. Je n'ai ni compilé ni exécuté l'application dans le cadre de la rédaction de ce rapport : les comportements décrits sont ceux que le code permet d'attendre, et non ceux d'une exécution observée.")
para("**Vérifications menées par l'étudiant pendant le stage :** [À COMPLÉTER : nature des tests manuels (outil utilisé), "
     "scénarios joués, résultats observés, captures d'écran éventuelles].")

# ----------------------------------------------------------------- 5. difficultés
h1("5. Difficultés, limites et analyse critique")
h2("5.1 Points techniques notables")
para("Les points ci-dessous sont des constats faits à la lecture du code et de l'historique Git. Ils ne sont pas présentés comme des difficultés vécues : les difficultés effectivement rencontrées et la manière dont je les ai résolues doivent être précisées par l'étudiant [À COMPLÉTER].", keep_next=True)
bullets([
    "**Exception sans message dans le gestionnaire global.** Le commit 4644688 remplace Map.of par un HashMap dans "
    "GlobalExceptionHandler. Map.of refuse les valeurs nulles : une exception sans message provoquait une "
    "NullPointerException à l'intérieur du gestionnaire.",
    "**Nom de table réservé.** L'entité utilisateur est mappée sur app_user, le commentaire de la classe signalant un "
    "conflit possible avec le mot réservé « user ».",
    "**Chargement paresseux et sérialisation.** Les associations ManyToOne sont chargées à la demande (LAZY) et open-in-view est désactivé : la conversion en DTO doit donc avoir lieu dans les services, pendant la session de persistance. Ce choix impose de rendre explicite l'accès aux associations.",
    "**Évolution de la règle de suppression.** Le commit 1db5da3 (15 juillet) ajoute à la suppression un contrôle des "
    "contrats liés, via existsByClientId. Le commit 4a0dfb8 (29 juillet) passe à l'archivage logique, et ce contrôle "
    "n'est plus appelé. Le dépôt ne documente pas la raison de ce changement.",
    "**Annotation sans effet.** Les @PreAuthorize figurent sans @EnableMethodSecurity : la protection effective est "
    "celle des chemins. Une annotation ne garantit donc pas à elle seule une restriction.",
    "**Requêtes inter-origines.** Le frontend et l'API tournent sur deux ports différents. La configuration CORS "
    "autorise l'origine http://localhost:4200.",
])
h2("5.2 Limites et pistes d'amélioration")
para("Le tableau {ref:tab_limites} synthétise les limites constatées dans le code et les pistes associées. Aucune de ces limites n'a été corrigée dans le cadre de ce rapport, qui décrit le projet tel qu'il est.", keep_next=True)
caption("Tableau", "tab_limites", "Limites constatées et pistes d'amélioration", above=True)
table(["Limite", "Constat dans le code", "Piste d'amélioration"], [
    ["Inscription publique avec rôle", "/api/auth/register est public et accepte le champ role : un appelant anonyme peut créer un ADMIN",
     "Réserver la création de comptes à un ADMIN authentifié ; ne pas accepter le rôle depuis le client"],
    ["Autorisation par méthode inopérante", "@PreAuthorize sans @EnableMethodSecurity ; un AGENT peut supprimer un client",
     "Activer la sécurité par méthode ; écrire des tests vérifiant les refus"],
    ["Suppression sans contrôle", "RM11 non appliquée", "Décider si l'archivage d'un client lié à un contrat est autorisé, puis l'appliquer ou le documenter"],
    ["Dates non vérifiées à la mise à jour", "updateContract ne contrôle pas endDate par rapport à startDate",
     "Appliquer la règle RM1 à la mise à jour"],
    ["Erreurs toutes en 400", "Gestionnaire unique pour RuntimeException, y compris « non trouvé »",
     "Gestionnaires par type d'exception (404, 403, 409) avec un format d'erreur normalisé"],
    ["Secrets en clair", "Clé JWT et identifiants PostgreSQL par défaut dans application.yaml et docker-compose.yml",
     "Variables d'environnement ; secrets hors du dépôt"],
    ["Justificatifs non contrôlés", "Pas de contrôle de type ni de taille ; pas de volume Docker dédié",
     "Liste des types autorisés, taille maximale, volume dédié"],
    ["Numéros générés par l'horloge", "Numéros construits avec les millisecondes modulo 100 000 : risque de collision",
     "Séquence en base de données ou identifiant UUID"],
    ["Schéma géré par Hibernate", "spring.jpa.hibernate.ddl-auto = update ; aucune migration", "Migrations versionnées (Flyway ou Liquibase)"],
    ["Tests insuffisants", "Un test unitaire ; aucun test de contrôleur ni d'intégration", "Tests MockMvc et tests d'intégration sur une base dédiée"],
    ["Justificatifs versionnés", "Le dossier uploads/ contient deux fichiers .jpg suivis par Git ; .gitignore ne l'exclut pas (contenu non examiné)", "Retirer ces fichiers du suivi et ignorer le dossier ; ne pas versionner de données personnelles"],
    
    ], [3.6, 6.2, 5.2], size=9.5)
h2("5.3 Analyse critique")
para("Le projet réussit sa partie structurante : les règles métier sont centralisées dans la couche service, les "
     "entités ne sont pas exposées directement grâce aux DTO, et la documentation est générée à partir du code. Il "
     "présente toutefois un écart entre ce qu'il affiche et ce qu'il garantit. Les annotations de rôle et l'absence "
     "de contrôle sur les dates de contrat en sont des exemples. Le plus grave est l'inscription publique avec choix "
     "du rôle, qui annule en pratique le contrôle d'accès par rôle.")


# ----------------------------------------------------------------- 6. bilan
h1("6. Bilan des compétences acquises")
h2("6.1 Compétences techniques")
para("Le tableau {ref:tab_comp} relie les compétences mobilisées aux éléments observables dans le dépôt. Le niveau "
     "proposé doit être validé par l'étudiant.", keep_next=True)
caption("Tableau", "tab_comp", "Compétences mobilisées et éléments observables", above=True)
table(["Domaine", "Mise en œuvre observée", "Niveau proposé (à valider)"], [
    ["Java 17 et programmation orientée objet", "Entités, énumérations, records, injection par constructeur", "Pratique"],
    ["API REST avec Spring Boot", "17 points d'accès, DTO, validation, gestion d'erreurs", "Pratique"],
    ["Persistance JPA et PostgreSQL", "Mapping, contraintes d'unicité et de non-nullité, relations, transactions", "Pratique"],
    ["Sécurité applicative", "JWT, BCrypt, filtre d'authentification, règles par chemin ; limites identifiées", "Initiation, à approfondir"],
    ["Documentation d'API", "springdoc-openapi, Swagger UI, schéma bearerAuth", "Pratique"],
    ["Conteneurisation", "Dockerfile en deux étapes, Docker Compose, contrôle de santé", "Initiation"],
    ["Angular et TypeScript", "Services, gardes, intercepteur, formulaires, envoi multipart", "Initiation"],
    ["Tests automatisés", "Un test unitaire avec JUnit 5 et Mockito", "Initiation"],
    ["Git et conventions de commit", "36 commits suivant la convention Conventional Commits", "Pratique"],
], [4.8, 7.0, 3.2])
para("Les compétences découvertes pendant le stage sont celles pour lesquelles le dépôt montre une première mise en "
     "œuvre : Angular, Docker et les tests. Les compétences approfondies sont celles qui apparaissent dans plusieurs "
     "commits successifs : API REST, persistance et validation.")
h2("6.2 Méthode de travail")
para("Le travail s'est appuyé sur une modélisation validée avec l'encadrante avant le développement, sur des commits "
     "de taille limitée aux messages normalisés, et sur une documentation mise à jour au fil de l'eau, dont le diagramme "
     "d'architecture modifié le 13 juillet. L'écart de méthode le plus net concerne les tests : ils sont peu nombreux, "
     "et une règle comme RM11 a pu être abandonnée sans qu'aucun test ne le signale.")
h2("6.3 Lien avec la formation")
para("Le stage a mobilisé la programmation orientée objet, les bases de données relationnelles et les notions de génie "
     "logiciel vues en première année : architecture en couches, séparation des responsabilités, tests. Il les a "
     "prolongées vers des usages concrets : sécurité des API par jetons, documentation OpenAPI et conteneurisation. "
     "[À CONFIRMER : correspondance exacte avec les modules du programme de FGL1].")
h2("6.4 Bilan personnel")
para("Ce stage m'a permis de passer d'une modélisation à une API complète, documentée et conteneurisée. L'analyse "
     "du projet m'a aussi appris qu'une annotation ne suffit pas à garantir une règle : il faut la vérifier par des "
     "tests qui échouent lorsqu'elle est violée. [Paragraphe personnel à compléter par l'étudiant : enseignements "
     "tirés des échanges avec l'encadrante, difficultés vécues, perspectives.]")

# ----------------------------------------------------------------- 7. conclusion
h1("7. Conclusion générale")
para("Le stage technicien effectué chez VERMEG du 1er au 31 juillet 2026 a abouti à une implémentation de mini-API "
     "REST de gestion d'assurance. Le dépôt contient un service Spring Boot qui expose 17 points d'accès pour les "
     "clients, les contrats, les sinistres et les utilisateurs. Ce service s'appuie sur PostgreSQL, sécurise l'accès "
     "par jetons JWT, applique dix règles métier côté serveur et publie sa documentation OpenAPI. Une interface Angular "
     "de 11 écrans consomme cette API, et l'ensemble est conteneurisé avec Docker Compose.")
para("Le bilan critique est tout aussi important. L'inscription publique avec choix du rôle, l'autorisation par "
     "méthode non activée, la règle de suppression abandonnée et la couverture de tests limitée sont les principaux "
     "points à corriger avant toute utilisation réelle. Le travail constitue une base technique que la suite du projet "
     "devra consolider.")
para("Le stage m'a permis de découvrir un environnement professionnel, de conduire une conception avec mon "
     "encadrante et d'acquérir des compétences concrètes en développement d'API.")

# ----------------------------------------------------------------- références
page_break()
h1("Références et webographie")
para("**Webographie** (sources consultées le 12 octobre 2026) :", keep_next=True)
for ref in [
    "[1] Nordic Capital, « Nordic Capital-backed Regnology acquires VERMEG's RegTech division – AGILE », communiqué de presse, novembre 2024. https://www.nordiccapital.com/news-views/press-releases/nordic-capital-backed-regnology-acquires-vermeg-s-regtech-division-agile-to-solidify-position-as-a-global-market-leader/",
    "[2] Crustdata, « Vermeg for Banking & Insurance Software – Headcount, Funding & Hiring Trends ». https://profiles.crustdata.com/company/vermeg-for-banking-insurance-software",
    "[3] RocketReach, « VERMEG for Banking & Insurance Software Information ». https://rocketreach.co/vermeg-for-banking-insurance-software-profile_b5c68ca2f42e0c9c",
    "[4] GitHub, dépôt « taheryhaia2/mini-api-assurance ». https://github.com/taheryhaia2/mini-api-assurance",
]:
    para(ref, space_after=3)
para("Les sources [1] à [3] sont des sources tierces, non officielles. [À COMPLÉTER : référence du site officiel de "
     "VERMEG, après consultation.] Le reste du rapport repose sur le code et l'historique du dépôt [4].")

# ----------------------------------------------------------------- annexes
h1("Annexe A — Points d'accès de l'API")
caption("Tableau", "tab_end", "Points d'accès de l'API, accès requis et code de succès", above=True)
table(["Méthode", "Chemin", "Accès requis", "Code"], [
    ["POST", "/api/auth/register", "Public (voir limite 5.2)", "201"],
    ["POST", "/api/auth/login", "Public", "200"],
    ["POST", "/api/clients", "ADMIN, AGENT", "201"],
    ["GET", "/api/clients", "ADMIN, AGENT", "200"],
    ["GET", "/api/clients/{id}", "ADMIN, AGENT", "200"],
    ["PUT", "/api/clients/{id}", "ADMIN, AGENT", "200"],
    ["DELETE", "/api/clients/{id}", "ADMIN, AGENT (@PreAuthorize inopérant)", "204"],
    ["POST", "/api/contracts", "ADMIN, AGENT", "201"],
    ["GET", "/api/contracts", "ADMIN, AGENT", "200"],
    ["GET", "/api/contracts/{id}", "ADMIN, AGENT", "200"],
    ["PUT", "/api/contracts/{id}", "ADMIN, AGENT (règle de chemin ; annotation ADMIN inopérante)", "200"],
    ["GET", "/api/contracts/client/{clientId}", "ADMIN, AGENT", "200"],
    ["POST", "/api/contracts/{contractId}/claims", "ADMIN, AGENT", "201"],
    ["GET", "/api/contracts/{contractId}/claims", "ADMIN, AGENT", "200"],
    ["GET", "/api/claims", "ADMIN, AGENT", "200"],
    ["GET", "/api/claims/{id}", "ADMIN, AGENT", "200"],
    ["GET", "/api/users", "ADMIN", "200"],
], [2.2, 6.0, 5.6, 1.7], size=9.5)

h1("Annexe B — Démarrage de l'application")
para("Les commandes ci-dessous sont celles décrites dans le dépôt. Elles n'ont pas été exécutées lors de la préparation "
     "de ce rapport.", keep_next=True)
bullets([
    "**Base de données :** créer une base PostgreSQL nommée assurance_db. Les identifiants se règlent par les "
    "variables SPRING_DATASOURCE_URL, SPRING_DATASOURCE_USERNAME et SPRING_DATASOURCE_PASSWORD.",
    "**Lancement local :** ./mvnw spring-boot:run, puis ouvrir http://localhost:8080/swagger-ui/index.html.",
    "**Lancement conteneurisé :** docker compose up --build.",
    "**Frontend :** dans le dossier frontend, npm install puis npm start ; l'interface est alors servie sur le port 4200.",
    "**Comptes de démonstration :** créés au démarrage de l'API. Leurs identifiants figurent dans le README du dépôt ; "
    "ils ne sont pas reproduits ici.",
])

doc.save(OUTPUT)
print("Rapport généré :", OUTPUT)
